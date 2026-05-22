from io import BytesIO
import unittest

from openpyxl import Workbook

from backend.parsing import parse_candidate_workbook


class CandidateWorkbookParsingTest(unittest.TestCase):
    def test_parse_valid_candidate_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["사번", "이름", "부서", "직위", "주재원 적합도", "사진URL", "발령내역"])
        sheet.append(["E24017", "김도현", "품질관리팀", "과장", 92, "https://example.com/e24017.jpg", "2024-01|품질관리팀|과장"])
        sheet.append(["E21884", "박준호", "영업팀", "차장", 95])

        result = parse_candidate_workbook(_to_bytes(workbook))

        self.assertEqual(len(result.candidates), 2)
        self.assertEqual(result.errors, [])
        self.assertEqual(result.candidates[0].data["employee_id"], "E24017")
        self.assertEqual(result.candidates[0].data["expat_fit"], 92)
        self.assertEqual(result.candidates[0].data["photo_url"], "https://example.com/e24017.jpg")
        self.assertEqual(result.candidates[0].data["work_history"], "2024-01|품질관리팀|과장")

    def test_missing_required_values_are_reported(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["사번", "이름"])
        sheet.append(["E24017", None])

        result = parse_candidate_workbook(_to_bytes(workbook))

        self.assertEqual(result.candidates, [])
        self.assertEqual(len(result.errors), 1)
        self.assertEqual(result.errors[0].row_number, 2)

    def test_language_score_header_maps_to_numeric_score(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["사번", "이름", "어학", "어학점수", "어학종류"])
        sheet.append(["E24017", "김도현", "TOEIC", 890, "TOEIC"])

        result = parse_candidate_workbook(_to_bytes(workbook))

        self.assertEqual(result.errors, [])
        self.assertEqual(result.candidates[0].data["language"], "TOEIC")
        self.assertEqual(result.candidates[0].data["language_score"], 890)


def _to_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


if __name__ == "__main__":
    unittest.main()
