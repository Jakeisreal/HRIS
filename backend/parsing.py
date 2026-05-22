from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any, BinaryIO

from openpyxl import load_workbook


COLUMN_ALIASES = {
    "employee_id": {"employee_id", "employee id", "id", "사번", "직원번호"},
    "name": {"name", "이름", "성명"},
    "dept": {"dept", "department", "부서", "소속"},
    "job_family": {"job_family", "job family", "직군", "직무군"},
    "position": {"position", "직위", "직급"},
    "grade": {"grade", "등급", "직급등급"},
    "hire_date": {"hire_date", "hire date", "입사일", "입사일자"},
    "location": {"location", "근무지", "사업장"},
    "email": {"email", "e-mail", "이메일", "메일"},
    "tenure": {"tenure", "근속", "근속년수"},
    "performance": {"performance", "평가", "성과평가"},
    "performance_2024": {"performance_2024", "2024평가", "2024 평가"},
    "performance_2025": {"performance_2025", "2025평가", "2025 평가"},
    "performance_2026": {"performance_2026", "2026평가", "2026 평가", "최근평가", "최근 평가"},
    "leadership": {"leadership", "리더십"},
    "language": {"language", "어학", "어학점수", "어학명"},
    "language_type": {"language_type", "language type", "어학종류", "시험종류"},
    "language_score": {"language_score", "language score", "어학숫자", "어학점수숫자"},
    "overseas": {"overseas", "해외경험", "주재경험"},
    "overseas_country": {"overseas_country", "overseas country", "해외국가", "주재국가"},
    "overseas_type": {"overseas_type", "overseas type", "해외유형", "주재유형"},
    "overseas_months": {"overseas_months", "overseas months", "해외개월", "해외경험개월"},
    "certificate": {"certificate", "자격", "자격증"},
    "expat_fit": {"expat_fit", "expat fit", "주재원적합도", "주재원 적합도"},
    "leader_fit": {"leader_fit", "leader fit", "팀장적합도", "팀장 적합도"},
    "purpose": {"purpose", "후보목적", "선정목적"},
}

REQUIRED_FIELDS = ("employee_id", "name")


@dataclass(frozen=True)
class ParsedCandidate:
    row_number: int
    data: dict[str, Any]


@dataclass(frozen=True)
class ParseError:
    row_number: int
    message: str


@dataclass(frozen=True)
class ParseResult:
    candidates: list[ParsedCandidate]
    errors: list[ParseError]
    headers: list[str]


def parse_candidate_workbook(file_obj: BinaryIO | BytesIO) -> ParseResult:
    if not isinstance(file_obj, BytesIO):
        file_obj = BytesIO(file_obj.read())
        file_obj.seek(0)

    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        return ParseResult(candidates=[], errors=[ParseError(1, "업로드 파일이 비어 있습니다.")], headers=[])

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    field_by_index = _map_headers(headers)
    missing = [field for field in REQUIRED_FIELDS if field not in field_by_index.values()]
    if missing:
        return ParseResult(
            candidates=[],
            errors=[ParseError(1, f"필수 컬럼이 없습니다: {', '.join(missing)}")],
            headers=headers,
        )

    candidates: list[ParsedCandidate] = []
    errors: list[ParseError] = []

    for offset, row in enumerate(rows, start=2):
        if not row or all(value in (None, "") for value in row):
            continue

        item: dict[str, Any] = {}
        for index, value in enumerate(row):
            field = field_by_index.get(index)
            if field:
                item[field] = _clean_cell(value)

        missing_values = [field for field in REQUIRED_FIELDS if not item.get(field)]
        if missing_values:
            errors.append(ParseError(offset, f"필수 값이 없습니다: {', '.join(missing_values)}"))
            continue

        item["employee_id"] = str(item["employee_id"]).strip()
        item["name"] = str(item["name"]).strip()
        candidates.append(ParsedCandidate(row_number=offset, data=item))

    return ParseResult(candidates=candidates, errors=errors, headers=headers)


def _map_headers(headers: list[str]) -> dict[int, str]:
    normalized_aliases = {
        field: {_normalize_header(alias) for alias in aliases}
        for field, aliases in COLUMN_ALIASES.items()
    }

    mapped: dict[int, str] = {}
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        for field, aliases in normalized_aliases.items():
            if normalized in aliases:
                mapped[index] = field
                break
    return mapped


def _normalize_header(value: str) -> str:
    return value.strip().lower().replace("_", "").replace(" ", "")


def _clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value
